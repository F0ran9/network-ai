# -*- coding: utf-8 -*-
"""报告导出器: 验收报告 / 割接方案, 支持 Markdown / DOCX / PDF 三种格式.

PDF 用 reportlab 生成, 注册 Windows 自带中文字体 (simhei/msyh/等).
"""
import io
import datetime

# ============================================================
# PDF 公共部分: 中文字体注册
# ============================================================
_PDF_FONT = "Helvetica"
_PDF_CN_FALLBACK_REGISTERED = False


def _register_pdf_cn_font():
    """注册 Windows 自带中文字体, 返回字体名; 全部失败则用 Helvetica."""
    global _PDF_FONT, _PDF_CN_FALLBACK_REGISTERED
    if _PDF_CN_FALLBACK_REGISTERED:
        return _PDF_FONT
    _PDF_CN_FALLBACK_REGISTERED = True
    import os
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [
        ("SimHei", r"C:\Windows\Fonts\simhei.ttf"),
        ("MSYaHei", r"C:\Windows\Fonts\msyh.ttc"),
        ("SimSun", r"C:\Windows\Fonts\simsun.ttc"),
        ("DengXian", r"C:\Windows\Fonts\Deng.ttf"),
    ]
    for name, path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path, subfontIndex=0))
                _PDF_FONT = name
                break
            except Exception:
                continue
    return _PDF_FONT


def _pdf_styles():
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    font = _register_pdf_cn_font()
    return {
        "title": ParagraphStyle("title", fontName=font, fontSize=20, alignment=TA_CENTER, spaceAfter=10),
        "h1": ParagraphStyle("h1", fontName=font, fontSize=14, spaceBefore=14, spaceAfter=6, textColor="#1a5fb4"),
        "body": ParagraphStyle("body", fontName=font, fontSize=10.5, leading=16, spaceAfter=6),
        "cell": ParagraphStyle("cell", fontName=font, fontSize=8.5, leading=11),
        "cellh": ParagraphStyle("cellh", fontName=font, fontSize=9, leading=12, fontNameBold=None),
        "code": ParagraphStyle("code", fontName="Courier", fontSize=8, leading=11),
    }


# ============================================================
# 验收报告
# ============================================================
def build_acceptance_report_md(d: dict) -> str:
    now = d.get("time") or datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    lines = [
        f"# 连通性验收报告",
        "",
        f"- **项目**: {d.get('projectName') or '网络变更'}",
        f"- **验收设备**: {d.get('host') or '-'} ({d.get('deviceType') or '-'})",
        f"- **源地址**: {d.get('source') or '-'}",
        f"- **执行人**: {d.get('operator') or '-'}",
        f"- **验收时间**: {now}",
        "",
        f"## 结果汇总",
        "",
        f"**{'✅ 全部通过' if d.get('allPass') else '⚠️ 存在不通项'}** — {d.get('passed', 0)}/{d.get('total', 0)} 通过",
        "",
        "| 目标 | IP | 结果 | 说明 |",
        "|---|---|---|---|",
    ]
    for r in d.get("results") or []:
        lines.append(f"| {r.get('name', '')} | {r.get('ip', '')} | "
                     f"{'PASS ✅' if r.get('pass') else 'FAIL ❌'} | {r.get('detail', '')} |")
    lines += ["", "---", "*本报告由网络拓扑配置平台自动生成*", ""]
    return "\n".join(lines)


def build_acceptance_report_docx(d: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from change_plan import _set_cn_font, _para, _table, _heading

    doc = Document()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("连通性验收报告"); _set_cn_font(run, size=20, bold=True)
    now = d.get("time") or datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"项目: {d.get('projectName') or '网络变更'}    验收时间: {now}")
    _set_cn_font(run, size=9, color=(0x99, 0x99, 0x99))

    _heading(doc, "一、基本信息", 1)
    _table(doc, ["验收设备", "设备类型", "源地址", "执行人"],
           [[d.get("host") or "-", d.get("deviceType") or "-", d.get("source") or "-", d.get("operator") or "-"]])

    _heading(doc, "二、结果汇总", 1)
    all_pass = d.get("allPass")
    _para(doc, f"{'全部通过 — 真的通 ✓' if all_pass else '存在不通项!'}  ({d.get('passed', 0)}/{d.get('total', 0)} 通过)",
          bold=True, color=(0x1a, 0x7f, 0x37) if all_pass else (0xcc, 0x33, 0x00))

    _heading(doc, "三、验收明细", 1)
    rows = [[r.get("name", ""), r.get("ip", ""),
             "PASS" if r.get("pass") else "FAIL", r.get("detail", "")]
            for r in d.get("results") or []]
    _table(doc, ["目标", "IP", "结果", "说明"], rows)

    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


def build_acceptance_report_pdf(d: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    s = _pdf_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm,
                            leftMargin=2.2 * cm, rightMargin=2.2 * cm)
    story = [
        Paragraph("连通性验收报告", s["title"]),
        Paragraph(f"项目: {d.get('projectName') or '网络变更'}    "
                  f"验收时间: {d.get('time') or datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                  ParagraphStyle("meta", parent=s["body"], fontSize=9, textColor="#999999")),
        Spacer(1, 8),
        Paragraph("一、基本信息", s["h1"]),
        Table([["验收设备", "设备类型", "源地址", "执行人"],
               [d.get("host") or "-", d.get("deviceType") or "-",
                d.get("source") or "-", d.get("operator") or "-"]],
              style=_pdf_table_style(len(d.get("results") or []) + 1, cols=4)),
        Paragraph("二、结果汇总", s["h1"]),
        Paragraph(("✅ 全部通过 — 真的通" if d.get("allPass") else "⚠️ 存在不通项!")
                  + f"  ({d.get('passed', 0)}/{d.get('total', 0)} 通过)", s["body"]),
        Paragraph("三、验收明细", s["h1"]),
    ]
    head = [Paragraph(x, s["cell"]) for x in ["目标", "IP", "结果", "说明"]]
    data = [head]
    for r in d.get("results") or []:
        data.append([Paragraph(str(r.get("name", "")), s["cell"]),
                     Paragraph(str(r.get("ip", "")), s["cell"]),
                     Paragraph("PASS" if r.get("pass") else "FAIL", s["cell"]),
                     Paragraph(str(r.get("detail", "")), s["cell"])])
    tbl = Table(data, colWidths=[4.5 * cm, 3.5 * cm, 1.8 * cm, 6.6 * cm], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef3fa")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafbfc")]),
    ]))
    story.append(tbl)
    doc.build(story)
    return buf.getvalue()


def _pdf_table_style(rows, cols):
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle
    return TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef3fa")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ])


# ============================================================
# 割接方案: MD / PDF (docx 已在 change_plan.py)
# ============================================================
ROLE_NAME = {"router": "路由器", "security": "安全设备", "core": "核心交换机",
             "aggregation": "汇聚交换机", "access": "接入交换机", "wireless": "无线",
             "endpoint": "终端"}

def _plan_devices(d):
    return d.get("devices") or []


def _plan_commands(dev):
    return [c for c in (dev.get("commands") or [])
            if c.strip() and not c.strip().startswith(("!", "#")) and "WARNING" not in c]


def build_change_plan_md(d: dict) -> str:
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    devices = _plan_devices(d)
    changed = [x for x in devices if _plan_commands(x)]
    lines = [
        f"# 网络变更割接方案",
        "",
        f"> **项目**: {d.get('projectName') or '网络变更项目'}  ",
        f"> **厂商**: {d.get('vendor') or '-'}  ",
        f"> **方案生成**: {now}（网络拓扑配置平台）",
        "",
        "## 一、变更概述",
        "",
        f"本次变更共涉及 {len(devices)} 台网络设备, 其中 {len(changed)} 台有配置变更内容。"
        "变更采用 dry-run 预览机制, 真实下发前自动备份 running-config, 异常可按回退方案恢复。",
        "",
        "## 二、设备清单",
        "",
        "| 序号 | 设备名称 | 角色 | 管理 IP | 变更命令数 |",
        "|---|---|---|---|---|",
    ]
    for i, dev in enumerate(devices):
        lines.append(f"| {i + 1} | {dev.get('name') or '-'} | {ROLE_NAME.get(dev.get('role', ''), '其他')} "
                     f"| {dev.get('host') or '未填'} | {len(_plan_commands(dev))} |")
    lines += ["", "## 三、配置变更明细", ""]
    if not changed:
        lines.append("本次变更无配置命令内容。")
    for dev in changed:
        lines.append(f"### {dev.get('name')}（{ROLE_NAME.get(dev.get('role', ''), '其他')} / {dev.get('host') or 'IP 未填'}）")
        lines.append("")
        lines.append("```")
        lines += _plan_commands(dev) or ["(无命令)"]
        lines.append("```")
        lines.append("")
    lines += [
        "## 四、回退方案",
        "",
        "1. 每台设备下发前自动备份 running-config, 备份文件即回退基线;",
        "2. 异常时在平台「备份回滚」页签选择变更前备份, 先 dry-run 预览再回滚;",
        "3. 回滚前平台会再次备份当前配置, 保证回退可回退;",
        "4. 回退后重新执行验收清单确认业务恢复;",
        "5. 建议变更窗口预留 30 分钟回退缓冲, 超时上报值班领导。",
        "",
        "## 五、连通性验收清单",
        "",
        "| 源设备 | 源地址 | 目标 | 目标 IP | 验收命令 | 预期 |",
        "|---|---|---|---|---|---|",
    ]
    n_acc = 0
    for acc in d.get("acceptance") or []:
        for t in acc.get("targets") or []:
            n_acc += 1
            if n_acc > 80: break
            lines.append(f"| {acc.get('source') or '-'} | {acc.get('sourceIp') or '-'} "
                         f"| {t.get('name') or '-'} | {t.get('ip') or '-'} "
                         f"| ping -a {acc.get('sourceIp') or '源IP'} {t.get('ip')} (5包) | PASS |")
    lines += [
        "",
        "## 六、签字确认",
        "",
        "| 变更执行人 | 复核人 | 开始时间 | 完成时间 | 结果 |",
        "|---|---|---|---|---|",
        f"| {d.get('operator') or ' '} | {d.get('reviewer') or ' '} |   |   |   |",
        "",
    ]
    return "\n".join(lines)


def build_change_plan_pdf(d: dict, topology_png: bytes = None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Image)
    from reportlab.lib.utils import ImageReader

    s = _pdf_styles()
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    devices = _plan_devices(d)
    changed = [x for x in devices if _plan_commands(x)]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm,
                            leftMargin=2.2 * cm, rightMargin=2.2 * cm)
    story = [
        Paragraph("网络变更割接方案", s["title"]),
        Paragraph(f"项目: {d.get('projectName') or '网络变更项目'}    厂商: {d.get('vendor') or '-'}    生成: {now}",
                  ParagraphStyle("meta", parent=s["body"], fontSize=9, textColor="#999999")),
        Paragraph("一、变更概述", s["h1"]),
        Paragraph(f"本次变更共涉及 {len(devices)} 台网络设备, 其中 {len(changed)} 台有配置变更内容。"
                  "变更采用 dry-run 预览机制, 真实下发前自动备份 running-config, 异常可按回退方案恢复。",
                  s["body"]),
        Paragraph("二、网络拓扑图", s["h1"]),
    ]
    if topology_png:
        try:
            img = ImageReader(io.BytesIO(topology_png))
            iw, ih = img.getSize()
            w = 16 * cm
            story.append(Image(io.BytesIO(topology_png), width=w, height=w * ih / iw))
        except Exception:
            story.append(Paragraph("(拓扑图插入失败, 请参考平台画布)", s["body"]))
    else:
        story.append(Paragraph("(未提供拓扑图, 请参考平台画布)", s["body"]))

    story.append(Paragraph("三、设备清单", s["h1"]))
    head = [Paragraph(x, s["cell"]) for x in ["序号", "设备名称", "角色", "管理 IP", "命令数"]]
    data = [head]
    for i, dev in enumerate(devices):
        data.append([Paragraph(str(i + 1), s["cell"]),
                     Paragraph(str(dev.get("name") or "-"), s["cell"]),
                     Paragraph(ROLE_NAME.get(dev.get("role", ""), "其他"), s["cell"]),
                     Paragraph(str(dev.get("host") or "未填"), s["cell"]),
                     Paragraph(str(len(_plan_commands(dev))), s["cell"])])
    tbl = Table(data, colWidths=[1.4 * cm, 5 * cm, 3 * cm, 4 * cm, 2 * cm], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef3fa")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(tbl)

    story.append(Paragraph("四、配置变更明细", s["h1"]))
    if not changed:
        story.append(Paragraph("本次变更无配置命令内容。", s["body"]))
    for dev in changed:
        story.append(Paragraph(f"{dev.get('name')}（{ROLE_NAME.get(dev.get('role', ''), '其他')}"
                               f" / {dev.get('host') or 'IP 未填'}）",
                               ParagraphStyle("dev", parent=s["body"], fontSize=11, fontNameBold=None)))
        code = Table([[Paragraph("<br/>".join(
            (c.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
             .replace(" ", "&nbsp;") for c in (_plan_commands(dev) or ["(无命令)"]))), s["code"])]],
            colWidths=[16 * cm])
        code.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f2f2f2")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#dddddd")),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(code)
        story.append(Spacer(1, 6))

    story += [
        Paragraph("五、回退方案", s["h1"]),
        Paragraph("1. 每台设备下发前自动备份 running-config, 备份文件即回退基线;<br/>"
                  "2. 异常时在平台「备份回滚」页签选择变更前备份, 先 dry-run 预览再回滚;<br/>"
                  "3. 回滚前平台会再次备份当前配置, 保证回退可回退;<br/>"
                  "4. 回退后重新执行验收清单确认业务恢复;<br/>"
                  "5. 建议变更窗口预留 30 分钟回退缓冲, 超时上报值班领导。", s["body"]),
        Paragraph("六、连通性验收清单", s["h1"]),
    ]
    data = [[Paragraph(x, s["cell"]) for x in ["源设备", "源地址", "目标", "目标 IP", "验收命令", "预期"]]]
    for acc in d.get("acceptance") or []:
        for t in acc.get("targets") or []:
            data.append([Paragraph(str(acc.get("source") or "-"), s["cell"]),
                         Paragraph(str(acc.get("sourceIp") or "-"), s["cell"]),
                         Paragraph(str(t.get("name") or "-"), s["cell"]),
                         Paragraph(str(t.get("ip") or "-"), s["cell"]),
                         Paragraph(f"ping -a {acc.get('sourceIp') or '源IP'} {t.get('ip')}", s["cell"]),
                         Paragraph("PASS", s["cell"])])
    tbl = Table(data, colWidths=[2.6 * cm, 2.6 * cm, 2.6 * cm, 2.6 * cm, 4.4 * cm, 1.2 * cm], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef3fa")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(tbl)
    story += [
        Paragraph("七、签字确认", s["h1"]),
        Table([["变更执行人", "复核人", "开始时间", "完成时间", "结果"],
               [d.get("operator") or "", d.get("reviewer") or "", "", "", ""]],
              colWidths=[3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm], rowHeights=[0.6 * cm, 1.2 * cm],
              style=_pdf_table_style(2, 5)),
    ]
    doc.build(story)
    return buf.getvalue()
