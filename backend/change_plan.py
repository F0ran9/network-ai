# -*- coding: utf-8 -*-
"""割接方案 Word 文档生成器.

把平台里的拓扑、设备清单、变更命令、回退与验收方案
组装成一份可直接打印签字的《网络变更割接方案》.docx
"""
import io
import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ROLE_NAME = {
    "router": "路由器", "security": "安全设备", "core": "核心交换机",
    "aggregation": "汇聚交换机", "access": "接入交换机", "wireless": "无线",
    "endpoint": "终端", "": "其他",
}


def _set_cn_font(run, name="微软雅黑", size=None, bold=None, color=None):
    run.font.name = "Calibri"  # 西文
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)  # 中文
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _heading(doc, text, level=1):
    h = doc.add_heading("", level=level)
    run = h.add_run(text)
    _set_cn_font(run, size=15 if level == 1 else 12.5, bold=True,
                 color=(0x1a, 0x5f, 0xb4) if level == 1 else (0x33, 0x33, 0x33))
    return h


def _para(doc, text, size=10.5, bold=False, color=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_cn_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _code_block(doc, commands):
    """命令块: 等宽字体 + 浅灰底纹(用单元格表格模拟, 保证打印可读)."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    # 底纹
    shd = cell._element.get_or_add_tcPr().makeelement(
        qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "F2F2F2"})
    cell._element.tcPr.append(shd)
    cell.text = ""
    first = True
    for line in commands:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(8.5)
        p.paragraph_format.space_after = Pt(0)


def _table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Light Grid Accent 1"
    for i, htext in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(htext)
        _set_cn_font(run, size=9.5, bold=True)
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            _set_cn_font(run, size=9)
    return tbl


def build_change_plan(data: dict, topology_png: bytes = None) -> bytes:
    """data 字段: projectName, vendor, devices[{name,role,host,commands}],
    acceptance[{source,sourceIp,targets[{name,ip}]}], operator, reviewer"""
    doc = Document()
    # 页边距
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)

    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    project = data.get("projectName") or "网络变更项目"
    vendor = data.get("vendor") or "-"

    # ---- 封面标题 ----
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("网络变更割接方案")
    _set_cn_font(run, size=22, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(project)
    _set_cn_font(run, size=13, color=(0x66, 0x66, 0x66))
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"设备厂商: {vendor}    方案生成: {now}    生成工具: 网络拓扑配置平台")
    _set_cn_font(run, size=9, color=(0x99, 0x99, 0x99))

    devices = data.get("devices") or []
    acceptance = data.get("acceptance") or []
    changed = [d for d in devices if d.get("commands")]

    # ---- 一、变更概述 ----
    _heading(doc, "一、变更概述", 1)
    _para(doc, f"本次变更共涉及 {len(devices)} 台网络设备, 其中 {len(changed)} 台有配置变更内容。"
               "变更过程采用 dry-run 预览机制: 每台设备真实下发前先连接并预览将执行的命令, "
               "下发动作会自动备份设备当前 running-config, 出现异常可按本方案回退章节立即恢复。"
               "变更完成后按验收方案逐项 ping 验证业务网段真实可达。")

    # ---- 二、网络拓扑 ----
    _heading(doc, "二、网络拓扑图", 1)
    if topology_png:
        try:
            doc.add_picture(io.BytesIO(topology_png), width=Cm(16.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            _para(doc, "(拓扑图插入失败, 请参考平台画布)", color=(0xcc, 0, 0))
    else:
        _para(doc, "(未提供拓扑图, 请参考平台画布)", color=(0x99, 0x99, 0x99))

    # ---- 三、设备清单 ----
    _heading(doc, "三、设备清单", 1)
    rows = [[i + 1, d.get("name") or "-", ROLE_NAME.get(d.get("role", ""), d.get("role") or "其他"),
             d.get("host") or "未填", len(d.get("commands") or [])]
            for i, d in enumerate(devices)]
    _table(doc, ["序号", "设备名称", "角色", "管理 IP", "变更命令数"], rows)

    # ---- 四、配置变更明细 ----
    _heading(doc, "四、配置变更明细", 1)
    if not changed:
        _para(doc, "本次变更无配置命令内容 (仅输出设备清单与拓扑备案)。", color=(0x99, 0x99, 0x99))
    for d in changed:
        _heading(doc, f"{d.get('name')}  ({ROLE_NAME.get(d.get('role', ''), '其他')}"
                      f" / {d.get('host') or 'IP 未填'})", 2)
        _para(doc, "通过 SSH 登录后进入系统视图, 依次执行以下命令 (注释行已过滤):", size=9.5,
              color=(0x66, 0x66, 0x66))
        cmds = [c for c in (d.get("commands") or [])
                if c.strip() and not c.strip().startswith(("!", "#")) and "WARNING" not in c]
        _code_block(doc, cmds or ["(无命令)"])

    # ---- 五、回退方案 ----
    _heading(doc, "五、回退方案", 1)
    _para(doc, "1. 每台设备真实下发前, 平台会自动将该设备当前 running-config 备份至服务器备份目录, 备份文件即回退基线;")
    _para(doc, "2. 若变更后业务异常, 在平台「真机下发 → 备份回滚」页签选择该设备变更前的备份文件, 先 dry-run 预览再执行回滚;")
    _para(doc, "3. 回滚动作执行前平台会再次备份当前配置, 保证回退动作本身可回退;")
    _para(doc, "4. 回退完成后重新执行第六章验收清单, 确认业务网段恢复可达;")
    _para(doc, "5. 回退判定时限: 建议变更窗口内预留 30 分钟作为回退缓冲时间, 超时未恢复则上报值班领导。")

    # ---- 六、验收方案 ----
    _heading(doc, "六、连通性验收清单", 1)
    _para(doc, "变更完成后, 从各设备以业务源地址 ping 以下目标, 全部 PASS 方可关闭变更单:", size=9.5,
          color=(0x66, 0x66, 0x66))
    acc_rows = []
    for acc in acceptance:
        for t in (acc.get("targets") or []):
            acc_rows.append([
                acc.get("source") or "-", acc.get("sourceIp") or "-",
                t.get("name") or "-", t.get("ip") or "-",
                f"ping -a {acc.get('sourceIp') or '源IP'} {t.get('ip')} (5包)",
                "PASS",
            ])
    if acc_rows:
        _table(doc, ["源设备", "源地址", "目标", "目标 IP", "验收命令", "预期"], acc_rows[:80])
        if len(acc_rows) > 80:
            _para(doc, f"(清单共 {len(acc_rows)} 项, 表格仅展示前 80 项, 其余请在平台「连通性验收」页签执行)",
                  size=8.5, color=(0x99, 0x99, 0x99))
    else:
        _para(doc, "(未生成验收清单: 请先为设备维护管理 IP)", color=(0x99, 0x99, 0x99))

    # ---- 七、签字栏 ----
    _heading(doc, "七、签字确认", 1)
    sign = _table(doc, ["变更执行人", "复核人", "变更开始时间", "变更完成时间", "结果"],
                  [[data.get("operator") or "", data.get("reviewer") or "", "", "", ""]])
    sign.rows[1].height = Cm(1.2)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
